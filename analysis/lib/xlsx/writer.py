import numpy as np
import pandas as pd
from openpyxl.styles import Border, Color, Font, Side
from openpyxl.utils.cell import get_column_letter

from analysis.lib.xlsx.style import (
    add_acres_percent_header,
    add_caption,
    add_good_condition,
    center_header_style,
    divider_col_style,
    left_header_style,
    value_style,
)


def write_excel(
    xlsx: pd.ExcelWriter,
    df: pd.DataFrame,
    sheet_name: str,
    caption: str | None = None,
    column_widths: list[float] | None = None,
    area_columns: list[int] | None = None,
    percent_columns: list[int] | None = None,
    good_condition_info: dict | None = None,
    breaks: list[int] | None = None,
    add_percent_divider=False,
):
    """Write a dataframe to xlsx.

    If caption is provided, it will be added on the first row and all columns
    for that row will be merged with it.

    Parameters
    ----------
    xlsx : ExcelWriter
    df : DataFrame
    sheet_name : str
    caption : str | None, optional
        if provided, will be formatted into the first cell of the sheet as "Table <x>: <caption>"
    column_widths : list[float] | None, optional
        if provided, column widths will be set to these values
    area_columns : list[int] | None, optional
        0-based indexes of area columns within the columns of the dataframe.
        If provided, these will be formatted as integer or floating point values.
    percent_columns : list[int] | None, optional
        0-based indexes of percent value columns within the columns of the dataframe.
        If provided, these will be formatted as percents.
    breaks : list[int] | None, optional
        0-based indexes of the first row of each logical grouping of rows
        (e.g., if there are multiple rows per analysis unit).
    add_percent_divider : bool, optional
        if both area_columns and percent_columns are provided and this is True,
        it will add a divider column between the acres and percents sections.
    """

    column_widths = column_widths or []
    area_columns = area_columns or []
    percent_columns = percent_columns or []
    breaks = breaks or []

    if add_percent_divider and not (column_widths and area_columns and percent_columns):
        raise ValueError("add_percent_divider=True requires column_widths, area_columns, and percent_columns")

    if column_widths and len(column_widths) != len(df.columns):
        raise ValueError("column_widths must be same length as number of columns")

    # reserve rows for caption, acres / percent header, and good condition header
    num_caption_rows = 0
    num_acres_percent_header_rows = 0
    num_good_condition_header_rows = 0

    if caption is not None:
        # add gap row after caption
        num_caption_rows = 2
    if add_percent_divider:
        num_acres_percent_header_rows = 1
    if good_condition_info is not None:
        num_good_condition_header_rows = 1

    data_start_row = num_caption_rows + num_acres_percent_header_rows + num_good_condition_header_rows

    df.to_excel(xlsx, sheet_name=sheet_name, startrow=data_start_row, index=False)
    ws = xlsx.sheets[sheet_name]

    if caption is not None:
        add_caption(ws, table_counter=len(xlsx.sheets), caption=caption)

    ### Add percent divider column
    percent_divider_col_index = -1
    if add_percent_divider:
        percent_divider_col_index = percent_columns[0]
        ws.insert_cols(idx=percent_divider_col_index + 1, amount=1)
        column_widths.insert(percent_divider_col_index, 4)

        # shift percent columns to the right
        percent_columns = (np.array(percent_columns) + 1).tolist()

    ### add acres / percent header
    if add_percent_divider:
        add_acres_percent_header(ws, num_caption_rows + 1, area_columns=area_columns, percent_columns=percent_columns)

    ### set column widths
    for i, width in enumerate(column_widths):
        letter = get_column_letter(i + 1)
        ws.column_dimensions[letter].width = width

    ### set cell styles
    for col_idx, col in enumerate(ws.columns):
        is_divider = col_idx == percent_divider_col_index

        if col_idx == 0:
            col[data_start_row].style = left_header_style
        elif is_divider:
            col[data_start_row].style = divider_col_style
        else:
            col[data_start_row].style = center_header_style

        for row_idx, cell in enumerate(col[data_start_row + 1 :]):
            if is_divider:
                cell.style = divider_col_style
            else:
                cell.style = value_style
                if breaks and row_idx + 1 in breaks:
                    cell.border = Border(bottom=Side(border_style="thin", color="000000"))

            value = cell.value
            is_int = isinstance(value, (float, int)) and int(value) == value
            is_link = isinstance(value, str) and value.startswith(("http://", "https://"))

            if col_idx in area_columns:
                if is_int:
                    cell.number_format = "#,##0"
                else:
                    cell.number_format = "#,##0.00"
            elif col_idx in percent_columns:
                if is_int:
                    cell.number_format = "0%"
                else:
                    cell.number_format = "0.00%"
            elif is_link:
                cell.hyperlink = cell.value
                cell.font = Font(color=Color(index=4))

    ### Add good condition header
    # NOTE: we do this after setting styling above since it restyles cells
    if good_condition_info is not None:
        add_good_condition(ws, num_caption_rows + num_acres_percent_header_rows + 1, info=good_condition_info)

    ### Set outer border
    if add_percent_divider:
        col = get_column_letter(area_columns[0] + 1)
        for row in range(num_caption_rows + 1, ws.max_row + 1):
            cell = ws[f"{col}{row}"]
            cell.border = Border(
                top=cell.border.top,
                left=Side(border_style="thin", color="000000"),
                bottom=cell.border.bottom,
                right=cell.border.right,
            )

        col = get_column_letter(ws.max_column)
        for row in range(num_caption_rows + 1, ws.max_row + 1):
            cell = ws[f"{col}{row}"]
            cell.border = Border(
                top=cell.border.top,
                left=cell.border.left,
                bottom=cell.border.bottom,
                right=Side(border_style="thin", color="000000"),
            )

    for col_idx in range(1, ws.max_column + 1):
        if col_idx - 1 == percent_divider_col_index:
            continue

        cell = ws[f"{get_column_letter(col_idx)}{ws.max_row}"]
        cell.border = Border(
            top=cell.border.top,
            left=cell.border.left,
            bottom=Side(border_style="medium", color="000000"),
            right=cell.border.right,
        )
